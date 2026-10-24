from __future__ import annotations

import argparse
import json
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from train_segmentation import UNet


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DEFAULT = ROOT / "Code" / "ieom_paper_format.docx"
OUTPUT_DEFAULT = ROOT / "docs" / "Bone_Fracture_IEOM2026_Full_Paper.docx"
MANUSCRIPT_SOURCE = ROOT / "docs" / "Bone_Fracture_WSSS_IEOM_Draft.docx"
TEST_SUMMARY = ROOT / "docs" / "outputs" / "classifier_test_summary.json"
DATA_FIGURE = ROOT / "docs" / "outputs" / "dataset_distribution.png"
CAM_INDEX = ROOT / "outputs" / "pseudo_masks" / "segmentation_train_index.csv"
SEGMENTATION_MANIFEST = ROOT / "outputs" / "data" / "segmentation_train_manifest.csv"
SEGMENTATION_HISTORY = ROOT / "outputs" / "segmenters" / "unet_128_s42" / "history.csv"
SEGMENTATION_CHECKPOINT = ROOT / "outputs" / "segmenters" / "unet_128_s42" / "best.pt"


def set_run(run, size=10, bold=False, italic=False, color="000000"):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(
    document,
    text="",
    *,
    size=10,
    bold=False,
    italic=False,
    alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
    before=0,
    after=7,
    keep=False,
):
    paragraph = document.add_paragraph()
    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = keep
    if text:
        set_run(paragraph.add_run(text), size=size, bold=bold, italic=italic)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = add_paragraph(
        document,
        text,
        size=12 if level == 1 else 11,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        before=8 if level == 1 else 4,
        after=4 if level == 1 else 3,
        keep=True,
    )
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_caption(document, text):
    return add_paragraph(
        document,
        text,
        size=9,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        before=3,
        after=8,
        keep=True,
    )


def set_cell_fill(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def add_table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    for column, width in enumerate(widths):
        table.columns[column].width = Inches(width)
    for column, text in enumerate(headers):
        cell = table.rows[0].cells[column]
        cell.width = Inches(widths[column])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, "1F4E78")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(3)
        set_run(paragraph.add_run(text), size=9, bold=True, color="FFFFFF")
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cell = cells[column]
            cell.width = Inches(widths[column])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_number % 2 == 1:
                set_cell_fill(cell, "F2F6FA")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(3)
            set_run(paragraph.add_run(str(value)), size=9)
    add_paragraph(document, "", after=2)
    return table


def remove_body_content(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def font(size, bold=False):
    name = "arialbd.ttf" if bold else "arial.ttf"
    path = Path("C:/Windows/Fonts") / name
    try:
        return ImageFont.truetype(str(path), size)
    except OSError:
        return ImageFont.load_default()


def make_metric_figure(metrics, destination):
    width, height = 1800, 820
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.text((75, 38), "Held-out classifier performance", fill="#172B4D", font=font(46, True))
    draw.text(
        (75, 100),
        "Point estimates and 95% exact-hash-group bootstrap intervals",
        fill="#52657A",
        font=font(28),
    )
    rows = [
        ("Accuracy", "accuracy"),
        ("Fracture precision", "precision_fractured"),
        ("Fracture recall", "recall_fractured"),
        ("Fracture F1", "f1_fractured"),
        ("Fracture ROC AUC", "auc_fractured"),
    ]
    x0, x1 = 650, 1300
    top, row_height = 230, 108
    for tick in range(0, 9):
        x = x0 + (x1 - x0) * tick / 8
        draw.line((x, top - 35, x, top + row_height * len(rows) - 30), fill="#DDE4EC", width=2)
        draw.text((x - 18, top + row_height * len(rows) - 16), f"{tick / 10:.1f}", fill="#52657A", font=font(20))
    colors = ["#2F6B9A", "#E69F24", "#388E6D", "#6C63A8", "#B45A55"]
    for index, (label, key) in enumerate(rows):
        center_y = top + index * row_height + 22
        value = metrics[key]
        low, high = metrics[f"{key}_ci95_low"], metrics[f"{key}_ci95_high"]
        px = lambda v: x0 + (x1 - x0) * v / 0.8
        draw.text((75, center_y - 18), label, fill="#172B4D", font=font(28, True))
        draw.line((px(low), center_y, px(high), center_y), fill=colors[index], width=8)
        draw.line((px(low), center_y - 16, px(low), center_y + 16), fill=colors[index], width=4)
        draw.line((px(high), center_y - 16, px(high), center_y + 16), fill=colors[index], width=4)
        draw.ellipse((px(value) - 12, center_y - 12, px(value) + 12, center_y + 12), fill=colors[index])
        draw.text((x1 + 30, center_y - 16), f"{value:.3f} [{low:.3f}, {high:.3f}]", fill="#263746", font=font(23))
    draw.text((x0, height - 38), "Metric value", fill="#52657A", font=font(22))
    image.save(destination)


def heatmap_overlay(gray_image, cam):
    gray = np.asarray(gray_image.convert("L").resize((400, 400), Image.Resampling.BILINEAR), dtype=np.float32)
    heat = np.asarray(
        Image.fromarray(np.uint8(np.clip(cam, 0, 1) * 255)).resize((400, 400), Image.Resampling.BILINEAR),
        dtype=np.float32,
    ) / 255
    red = np.clip(255 * (0.35 + 0.65 * heat), 0, 255)
    green = np.clip(110 * (1 - heat), 0, 255)
    blue = np.clip(35 * (1 - heat), 0, 255)
    base = np.stack((gray, gray, gray), axis=-1)
    color = np.stack((red, green, blue), axis=-1)
    mixed = np.uint8(np.clip(0.55 * base + 0.45 * color, 0, 255))
    return Image.fromarray(mixed, mode="RGB")


def make_qualitative_figure(destination):
    if not all(path.exists() for path in (CAM_INDEX, SEGMENTATION_CHECKPOINT)):
        raise FileNotFoundError("CAM index or trained U-Net checkpoint is missing.")
    index = pd.read_csv(CAM_INDEX)
    examples = index.loc[index["split"].eq("train")].sort_values("relative_path")
    if examples.empty:
        raise ValueError("No training-only CAM examples are available.")
    row = examples.iloc[0]
    image_path, cam_path, mask_path = map(Path, (row.image_path, row.cam_path, row.mask_path))
    for path in (image_path, cam_path, mask_path):
        if not path.is_file():
            raise FileNotFoundError(f"Training example artifact is missing: {path}")

    with Image.open(image_path) as source:
        source = source.convert("L")
    cam = np.load(cam_path)
    if cam.ndim != 2:
        raise ValueError(f"Expected a 2-D CAM, got shape {cam.shape}.")
    with Image.open(mask_path) as mask_file:
        mask = mask_file.convert("L")

    model = UNet()
    with torch.serialization.safe_globals(
        [np._core.multiarray.scalar, np.dtype, type(np.dtype(np.float64))]
    ):
        checkpoint = torch.load(SEGMENTATION_CHECKPOINT, map_location="cpu", weights_only=True)
    model.load_state_dict(checkpoint["model_state_dict"])
    model.eval()
    resized = source.resize((128, 128), Image.Resampling.BILINEAR)
    tensor = torch.from_numpy(np.asarray(resized, dtype=np.float32).copy() / 255.0)[None, None]
    with torch.no_grad():
        prediction = (torch.sigmoid(model(tensor))[0, 0].numpy() >= 0.5).astype(np.uint8) * 255

    panels = [
        ("Training radiograph", source.convert("RGB")),
        ("Grad-CAM overlay", heatmap_overlay(source, cam)),
        ("CAM pseudo-mask (threshold 0.35)", mask),
        ("U-Net output (threshold 0.5)", Image.fromarray(prediction, mode="L")),
    ]
    panel_width, panel_height, heading_height = 430, 430, 76
    margin, gap = 42, 28
    width = margin * 2 + len(panels) * panel_width + (len(panels) - 1) * gap
    height = heading_height + panel_height + 70
    canvas = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(canvas)
    for i, (label, panel) in enumerate(panels):
        left = margin + i * (panel_width + gap)
        draw.text((left + 4, 18), label, fill="#172B4D", font=font(22, True))
        panel = panel.resize((panel_width, panel_height), Image.Resampling.NEAREST if "mask" in label.lower() or "output" in label.lower() else Image.Resampling.BILINEAR)
        if panel.mode != "RGB":
            panel = panel.convert("RGB")
        canvas.paste(panel, (left, heading_height))
        draw.rectangle((left, heading_height, left + panel_width - 1, heading_height + panel_height - 1), outline="#D1D9E0", width=2)
    canvas.save(destination)


def get_references():
    wanted = (
        "Abedeen,",
        "Efron,",
        "Geirhos,",
        "He,",
        "Nagy,",
        "Ronneberger,",
        "Selvaraju,",
    )
    source = Document(MANUSCRIPT_SOURCE)
    after_references = False
    references = []
    for paragraph in source.paragraphs:
        text = " ".join(paragraph.text.split())
        if text == "References":
            after_references = True
            continue
        if text == "Biographies":
            break
        if after_references and text.startswith(wanted):
            references.append(text)
    if len(references) != len(wanted):
        raise ValueError(f"Expected {len(wanted)} relevant references, found {len(references)}.")
    return references


def add_picture(document, path, width=6.3):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def build(template, output):
    required = (
        template,
        MANUSCRIPT_SOURCE,
        TEST_SUMMARY,
        DATA_FIGURE,
        CAM_INDEX,
        SEGMENTATION_MANIFEST,
        SEGMENTATION_HISTORY,
        SEGMENTATION_CHECKPOINT,
    )
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Required paper inputs are missing:\n" + "\n".join(missing))

    results = json.loads(TEST_SUMMARY.read_text(encoding="utf-8"))
    metrics = results["metrics"]
    split_counts = pd.read_csv(SEGMENTATION_MANIFEST)["split"].value_counts().to_dict()
    segmentation_history = pd.read_csv(SEGMENTATION_HISTORY)
    best_segmentation = segmentation_history.loc[
        segmentation_history["val_dice_against_training_labels"].idxmax()
    ]
    references = get_references()
    document = Document(str(template))
    remove_body_content(document)

    add_paragraph(
        document,
        "Proceedings of the International Conference on Industrial Engineering and Operations Management",
        size=12,
        bold=True,
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
    )
    add_paragraph(
        document,
        "A Leakage-Aware Baseline for Bone Fracture Classification and CAM-Derived Pseudo-Mask Segmentation",
        size=18,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=12,
        keep=True,
    )
    add_paragraph(document, "Shibli Sanjid Faheem", size=12, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_paragraph(document, "Department of Computer Science and Engineering", size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_paragraph(document, "Independent University, Bangladesh", size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_paragraph(document, "shiblisanjidfaheem@gmail.com", size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=1)
    add_paragraph(document, "Yousuf Abdullah", size=12, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_paragraph(document, "Department of Computer Science and Engineering", size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_paragraph(document, "Independent University, Bangladesh", size=12, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=9)

    add_paragraph(document, "Abstract", size=12, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    abstract = (
        "This study follows a leakage-aware workflow for bone-fracture image classification and weakly supervised "
        "pseudo-mask segmentation using the single image-folder dataset available in the project repository. The audit "
        "identified 1,927 images and 1,897 unique exact-image hashes; patient identifiers and expert pixel masks were "
        "unavailable. An ImageNet-pretrained ResNet-50 was selected using validation AUC and evaluated once on the held-out "
        f"test split (n={results['n']}). Test accuracy was {metrics['accuracy']:.3f} "
        f"(95% exact-hash-group bootstrap interval {metrics['accuracy_ci95_low']:.3f}-{metrics['accuracy_ci95_high']:.3f}), "
        f"and fracture-class AUC was {metrics['auc_fractured']:.3f} "
        f"({metrics['auc_fractured_ci95_low']:.3f}-{metrics['auc_fractured_ci95_high']:.3f}). Grad-CAM maps were generated "
        "only from training images and thresholded to produce pseudo-labels. A compact U-Net was trained on 734 training "
        "pairs and selected using 184 validation pairs; its best validation Dice and IoU against held-out pseudo-masks "
        f"were {best_segmentation['val_dice_against_training_labels']:.3f} and "
        f"{best_segmentation['val_iou_against_training_labels']:.3f}. These are pseudo-label agreement measures, not "
        "expert-mask segmentation results. The near-chance classification AUC and lack of expert annotations limit "
        "interpretation; no clinical or localization-performance claim is made."
    )
    add_paragraph(document, abstract, after=10)
    add_paragraph(document, "Keywords", size=12, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, after=1)
    add_paragraph(
        document,
        "Bone fracture, radiographs, ResNet-50, Grad-CAM, pseudo-mask, U-Net, weak supervision",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        after=8,
    )

    add_heading(document, "1. Introduction")
    add_paragraph(
        document,
        "Automated fracture analysis from radiographs is often studied first as image-level classification, while "
        "localization requires more detailed annotation. Image-folder labels are comparatively simple to collect, but "
        "they indicate whether an image is labeled fractured rather than where a fracture is located. A classifier can "
        "therefore produce a class prediction without providing an anatomically validated mask. Public resources such as "
        "FracAtlas and GRAZPEDWRI-DX illustrate the value of datasets with localization or expert annotations, while also "
        "showing why source and population differences must be considered (Abedeen et al. 2023; Nagy et al. 2022)."
    )
    add_paragraph(
        document,
        "Class activation maps can provide spatial evidence associated with a classifier decision (Selvaraju et al. 2017). "
        "However, a discriminative response map is not equivalent to a complete fracture annotation: its resolution, "
        "threshold, and classifier errors can all affect the apparent region. Training a segmentation network on such "
        "thresholded maps measures its ability to reproduce the generated targets, not its agreement with expert anatomy."
    )
    add_paragraph(
        document,
        "This paper reports the completed repository workflow as a preliminary empirical baseline. It separates the "
        "held-out classification evaluation from the train-only CAM and U-Net experiment, reports uncertainty for "
        "classification at the exact-image-hash group level, and treats pseudo-mask metrics as internal agreement only."
    )
    add_heading(document, "1.1 Objectives", level=2)
    add_paragraph(document, "The study has four objectives:", after=3)
    for objective in (
        "to audit the available image folders, exact duplicates, and presence of candidate expert masks;",
        "to train an ImageNet-initialized ResNet-50 and report held-out classification performance with group-bootstrap intervals;",
        "to generate Grad-CAM pseudo-masks exclusively from training images and train a compact U-Net without accessing the test split; and",
        "to distinguish pseudo-label agreement from expert-mask segmentation performance and identify the evidence still needed."
    ):
        paragraph = document.add_paragraph(style="List Paragraph")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.space_after = Pt(3)
        set_run(paragraph.add_run(objective), size=10)

    add_heading(document, "2. Literature Review")
    add_paragraph(
        document,
        "Residual networks provide a widely used transfer-learning backbone for image classification (He et al. 2016). "
        "Gradient-weighted class activation mapping uses gradients at a convolutional layer to form a class-specific "
        "localization map (Selvaraju et al. 2017). Such maps can be useful for visualization and weak supervision, but "
        "they should not be treated as expert labels without independent validation."
    )
    add_paragraph(
        document,
        "U-Net is an encoder-decoder architecture developed for biomedical image segmentation (Ronneberger et al. 2015). "
        "Its performance depends on the target masks used for training and evaluation. When CAM-derived targets are "
        "used, the resulting validation score describes agreement with those targets, and inherits their omissions and "
        "false activations."
    )
    add_paragraph(
        document,
        "Shortcut learning is a general risk when predictive image cues do not represent the intended clinical feature "
        "(Geirhos et al. 2020). Exact-hash grouping can prevent literal duplicate leakage but does not identify related "
        "studies or patients. The present experiment therefore uses hash grouping as a limited safeguard rather than a "
        "substitute for patient-level splitting."
    )
    add_caption(document, "Table 1: Positioning of the present workflow against relevant model roles")
    add_table(
        document,
        ["Approach", "Supervision or output", "Evaluation scope in this study"],
        [
            ("ResNet-50 (He et al. 2016)", "Image-level class labels", "Held-out test classification"),
            ("Grad-CAM (Selvaraju et al. 2017)", "Class-conditioned activation map", "Training-only pseudo-mask generation"),
            ("U-Net (Ronneberger et al. 2015)", "Dense segmentation prediction", "Validation agreement with pseudo-masks"),
            ("This workflow", "Image labels to CAM pseudo-labels to U-Net", "No expert-mask localization benchmark"),
        ],
        [1.55, 2.25, 2.7],
    )

    add_heading(document, "3. Methods")
    add_paragraph(
        document,
        "The workflow comprises data inventory and exact-hash checks, image-level classifier training, training-only "
        "Grad-CAM generation, threshold-based pseudo-mask creation, and U-Net training. The classification test split "
        "was used only for the classifier evaluation. It was excluded from CAM target generation and segmentation training."
    )
    add_heading(document, "3.1 Data inventory and leakage-aware preparation", level=2)
    add_paragraph(
        document,
        "The dataset is organized into train, validation, and test directories with fractured and not-fractured labels. "
        "The manifest contains 1,927 image files and 1,897 unique SHA-256 image hashes. The original folder split contains "
        "936 train, 600 validation, and 391 test files. The audit found no exact-hash groups crossing the original splits "
        "and no conflicting labels within duplicate groups. No candidate expert pixel masks, patient identifiers, or "
        "study identifiers were available. Exact hashes cannot rule out transformed near-duplicates or repeated patients."
    )
    add_heading(document, "3.2 Classifier training", level=2)
    add_paragraph(
        document,
        "An ImageNet-pretrained ResNet-50 with a new two-class head was trained using the image-level labels. Grayscale "
        "radiographs were converted to three channels, resized to 224 by 224 pixels, and normalized with ImageNet "
        "statistics. Training used class-balanced sampling, AdamW, a batch size of 16, head learning rate 1e-3, backbone "
        "learning rate 1e-5, and weight decay 1e-4. The classifier used seed 0, two head-warmup epochs, and ten total "
        "epochs. The checkpoint was selected by validation fracture-class AUC; epoch 5 was selected. The test split was "
        "not used for checkpoint selection or threshold tuning."
    )
    add_heading(document, "3.3 CAM generation and pseudo-mask construction", level=2)
    add_paragraph(
        document,
        "Grad-CAM maps were generated for the 918 unique training images only. A global threshold of 0.35 converted "
        "each map into a binary pseudo-mask. No validation or test image was used to create a training target. These "
        "maps are classifier-derived noisy labels and are not expert delineations."
    )
    add_heading(document, "3.4 U-Net training", level=2)
    add_paragraph(
        document,
        "The compact U-Net was trained on 734 training image-pseudo-mask pairs and validated on 184 separate pairs. "
        "The split was formed by exact-image-hash groups from the original training partition. Inputs were grayscale "
        "images resized to 128 by 128 pixels; training used batch size 8 and seed 42 for three epochs. The selected "
        "checkpoint was the epoch with highest validation Dice against the held-out pseudo-masks (epoch 2). The "
        "repository test split was not used in this segmentation experiment."
    )
    add_heading(document, "3.5 Metrics and uncertainty", level=2)
    add_paragraph(
        document,
        "Fractured was treated as the positive class at a fixed classification probability threshold of 0.5. Accuracy, "
        "fracture precision, recall, F1, and ROC AUC are reported. Percentile 95% confidence intervals use 2,000 "
        "bootstrap resamples of exact-image-hash groups with seed 42 (Efron 1979). These intervals quantify image-group "
        "sampling uncertainty, not patient-level uncertainty. U-Net Dice and IoU use a 0.5 output threshold and compare "
        "predictions with pseudo-masks only; they are not estimates of expert-mask accuracy."
    )

    add_heading(document, "4. Data Collection and Experimental Setup")
    add_paragraph(
        document,
        "The experiment used the single image-folder dataset present in the repository checkout. Its provenance, "
        "collection protocol, license, and patient grouping were not documented in the available project files. The "
        "dataset inventory is summarized in Figure 1. The classifier test set contains 391 images (193 fractured and "
        "198 not fractured). The classifier ran on CPU using Python 3.14.8, PyTorch 2.14.1, and torchvision 0.29.1. "
        "The U-Net experiment used the same CPU environment. No external dataset or expert-mask subset was added."
    )
    add_picture(document, DATA_FIGURE, width=6.25)
    add_caption(document, "Figure 1: Image counts by existing split and class; inventory counts are not model results.")

    add_heading(document, "5. Results and Discussion")
    add_heading(document, "5.1 Held-out classification results", level=2)
    add_paragraph(
        document,
        f"The best checkpoint was selected at epoch {results['checkpoint_epoch']} with validation fracture-class "
        f"AUC {results['checkpoint_validation_auc_fractured']:.3f}. On the held-out test split (n={results['n']}), "
        "accuracy was close to 0.5 and the ROC AUC was 0.508. The AUC 95% group-bootstrap interval includes 0.5, "
        "so this single experiment does not show reliable discrimination above chance. The fixed-threshold recall "
        "was higher than precision; these results are preliminary and do not establish clinical performance."
    )
    add_caption(document, "Table 2: Held-out classifier performance with 95% exact-hash-group bootstrap intervals")
    add_table(
        document,
        ["Metric", "Estimate", "95% bootstrap interval"],
        [
            ("Accuracy", f"{metrics['accuracy']:.3f}", f"[{metrics['accuracy_ci95_low']:.3f}, {metrics['accuracy_ci95_high']:.3f}]"),
            ("Fracture precision", f"{metrics['precision_fractured']:.3f}", f"[{metrics['precision_fractured_ci95_low']:.3f}, {metrics['precision_fractured_ci95_high']:.3f}]"),
            ("Fracture recall", f"{metrics['recall_fractured']:.3f}", f"[{metrics['recall_fractured_ci95_low']:.3f}, {metrics['recall_fractured_ci95_high']:.3f}]"),
            ("Fracture F1", f"{metrics['f1_fractured']:.3f}", f"[{metrics['f1_fractured_ci95_low']:.3f}, {metrics['f1_fractured_ci95_high']:.3f}]"),
            ("Fracture ROC AUC", f"{metrics['auc_fractured']:.3f}", f"[{metrics['auc_fractured_ci95_low']:.3f}, {metrics['auc_fractured_ci95_high']:.3f}]"),
        ],
        [1.75, 1.0, 3.75],
    )
    with tempfile.TemporaryDirectory(prefix="bone-fracture-paper-") as temporary:
        metric_figure = Path(temporary) / "test_metrics.png"
        qualitative_figure = Path(temporary) / "pseudo_mask_example.png"
        make_metric_figure(metrics, metric_figure)
        add_picture(document, metric_figure, width=6.25)
        add_caption(
            document,
            "Figure 2: Test classification estimates and 95% exact-hash-group bootstrap intervals; the AUC interval includes 0.5.",
        )

        add_heading(document, "5.2 CAM pseudo-mask and U-Net results", level=2)
        add_paragraph(
            document,
            f"The U-Net's best validation Dice was {best_segmentation['val_dice_against_training_labels']:.3f} and "
            f"IoU was {best_segmentation['val_iou_against_training_labels']:.3f}, both measured against generated "
            "pseudo-masks. These values describe agreement with the training target-generation process, not the "
            "correctness of the predicted fracture boundaries."
        )
        add_caption(document, "Table 3: Best U-Net validation results against held-out pseudo-masks")
        add_table(
            document,
            ["Validation measure", "Value", "Interpretation"],
            [
                ("Dice", f"{best_segmentation['val_dice_against_training_labels']:.3f}", "Pseudo-label agreement only"),
                ("IoU", f"{best_segmentation['val_iou_against_training_labels']:.3f}", "Pseudo-label agreement only"),
            ],
            [1.5, 0.8, 4.2],
        )
        make_qualitative_figure(qualitative_figure)
        add_picture(document, qualitative_figure, width=6.25)
        add_caption(
            document,
            "Figure 3: Illustrative training-split radiograph, Grad-CAM, thresholded pseudo-mask, and U-Net output. "
            "This example is not expert-annotated and is not a test image.",
        )
    add_paragraph(
        document,
        "Figure 3 is included to show the generated artifacts and their relationship, not to demonstrate successful "
        "localization. Its example is selected deterministically from the training index; no claim is made that this "
        "case is representative of the full dataset."
    )
    add_heading(document, "5.3 Interpretation and limitations", level=2)
    add_paragraph(
        document,
        "The classification AUC of 0.508 is near chance and its confidence interval spans 0.5. The experiment therefore "
        "does not establish a useful classifier, despite recall of 0.637 at the fixed threshold. The uncertainty "
        "intervals are based on exact-image groups rather than patients because patient metadata are unavailable. "
        "Additionally, the image-statistics diagnostic reported in the project audit is not proof that the model is "
        "free of shortcut learning (Geirhos et al. 2020)."
    )
    add_paragraph(
        document,
        "The U-Net inherits the quality and omissions of its CAM targets. Its validation set is separated by exact "
        "hash from the training pseudo-mask pairs, but both sets come from the same weak-label generation procedure. "
        "Without expert masks, the Dice and IoU values cannot be interpreted as anatomical segmentation performance. "
        "No segmentation test-set result is reported."
    )
    add_heading(document, "6. Conclusion")
    add_paragraph(
        document,
        "This paper documents a single-source, leakage-aware classification and pseudo-mask segmentation workflow "
        "using the repository dataset. The ResNet-50 achieved test accuracy 0.506 and AUC 0.508, with the AUC interval "
        "including chance-level performance. The U-Net achieved validation Dice 0.378 and IoU 0.267 against held-out "
        "CAM pseudo-masks, not expert annotations. These results are insufficient to support clinical or fracture "
        "localization claims. A verified dataset citation and license, patient- or study-level splits, repeated "
        "experiments, external validation, and expert pixel masks are needed for stronger conclusions."
    )
    add_heading(document, "Acknowledgements")
    add_paragraph(
        document,
        "No external funding or additional acknowledgements were provided for this draft.",
    )
    add_paragraph(
        document,
        "Declaration of generative AI use: AI-assisted tools were used to help organize and edit this manuscript. "
        "The reported results were taken from the project artifacts, and the authors remain responsible for checking "
        "the data, citations, interpretation, and final submission.",
    )
    add_heading(document, "References")
    for reference in references:
        paragraph = add_paragraph(document, reference, size=9, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4)
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)

    add_heading(document, "Biographies")
    add_paragraph(
        document,
        "Shibli Sanjid Faheem is an undergraduate student in Computer Science and Engineering with a minor in "
        "Management Information Systems at Independent University, Bangladesh.",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_paragraph(
        document,
        "Yousuf Abdullah is a co-author of this study and is affiliated with the Department of Computer Science and "
        "Engineering, Independent University, Bangladesh.",
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )

    used_images = {
        blip.get(qn("r:embed"))
        for blip in document.element.body.iter(qn("a:blip"))
    }
    for relationship_id, relationship in list(document.part.rels.items()):
        if relationship.reltype.endswith("/image") and relationship_id not in used_images:
            document.part.drop_rel(relationship_id)

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Created {output}")


def main():
    parser = argparse.ArgumentParser(description="Create a new IEOM-style bone-fracture paper from the supplied Word template.")
    parser.add_argument("--template", type=Path, default=TEMPLATE_DEFAULT)
    parser.add_argument("--output", type=Path, default=OUTPUT_DEFAULT)
    args = parser.parse_args()
    build(args.template, args.output)


if __name__ == "__main__":
    main()
